"""Text extraction helpers for company knowledge documents."""

import re
import unicodedata
from io import BytesIO

from docx import Document
from pypdf import PdfReader


def normalize_text(text):
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x00", "")

    lines = [
        re.sub(r"[^\S\n]+", " ", line).strip()
        for line in text.split("\n")
    ]

    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def extract_text(content, extension):
    if extension == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("Text files must use UTF-8 encoding.") from error

    elif extension == ".pdf":
        reader = PdfReader(BytesIO(content))

        if reader.is_encrypted:
            raise ValueError("Encrypted PDF files are not supported.")

        text = "\n\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    elif extension == ".docx":
        document = Document(BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        text = "\n\n".join(parts)

    else:
        raise ValueError("Supported file types are PDF, DOCX and TXT.")

    text = normalize_text(text)

    if not text:
        raise ValueError("The document contains no extractable text.")

    return text


def split_text(text, chunk_size=1200, overlap=200):
    if chunk_size <= 0 or not 0 <= overlap < chunk_size:
        raise ValueError("Overlap must be non-negative and smaller than chunk size.")

    text = normalize_text(text)
    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end == len(text):
            break

        start = end - overlap

    return chunks
